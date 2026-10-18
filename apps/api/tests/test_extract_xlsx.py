"""Spreadsheet text for search: every data row carries its column names."""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook

from quberi_api.services.rag.chunk import chunk_text
from quberi_api.services.rag.extract import extract_text

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _extract(wb: Workbook) -> str:
    buf = io.BytesIO()
    wb.save(buf)
    return extract_text(blob=buf.getvalue(), filename="t.xlsx", mime_type=XLSX).text


def test_each_data_row_is_written_with_its_column_names():
    wb = Workbook()
    ws = wb.active
    ws.title = "报价"
    ws.append(["零件", "单价", "数量"])
    ws.append(["螺丝 M6", 0.35, 1000])
    ws.append(["垫片", 0.1, None])  # an empty cell is left out, not printed as a blank
    text = _extract(wb)
    assert text.splitlines() == [
        "[Sheet: 报价]",
        "零件 | 单价 | 数量",
        "零件: 螺丝 M6 | 单价: 0.35 | 数量: 1000",
        "零件: 垫片 | 单价: 0.1",
    ]


def test_a_piece_from_the_middle_of_a_big_sheet_still_says_what_each_number_is():
    wb = Workbook()
    ws = wb.active
    ws.append(["零件", "单价", "数量"])
    for i in range(400):
        ws.append([f"part-{i:03d}", i, i * 2])
    chunks = chunk_text(_extract(wb), max_chars=600, overlap_chars=50)
    assert len(chunks) > 3
    middle = chunks[len(chunks) // 2].text
    assert "part-" in middle and "单价:" in middle and "数量:" in middle


def test_a_sheet_without_a_header_row_keeps_the_plain_form():
    wb = Workbook()
    ws = wb.active
    ws.append([1, 2, 3])
    ws.append([4, 5, 6])
    assert _extract(wb).splitlines() == ["[Sheet: Sheet]", "1 | 2 | 3", "4 | 5 | 6"]


def test_text_rows_that_are_not_a_header_are_left_alone():
    """First row with a number or a date in it is data, and so is a lone label."""

    wb = Workbook()
    ws = wb.active
    ws.append(["Q3", 2026, date(2026, 9, 30)])
    ws.append(["a", "b", "c"])
    first = _extract(wb).splitlines()[1]
    assert first.startswith("Q3 | 2026 | 2026-09-30")
    wb2 = Workbook()
    wb2.active.append(["just a title"])
    wb2.active.append(["body"])
    assert _extract(wb2).splitlines()[1:] == ["just a title", "body"]


def test_each_sheet_gets_its_own_header():
    wb = Workbook()
    a = wb.active
    a.title = "A"
    a.append(["name", "age"])
    a.append(["Ann", 30])
    b = wb.create_sheet("B")
    b.append(["city", "population"])
    b.append(["Oslo", 700000])
    text = _extract(wb)
    assert "name: Ann | age: 30" in text and "city: Oslo | population: 700000" in text
